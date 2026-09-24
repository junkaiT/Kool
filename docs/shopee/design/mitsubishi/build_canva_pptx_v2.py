"""Systems 1-3 covers in the System 4 layout the user built in Canva.

1080x1080. Logo top-left, StarMEX centred, that system's own NEA label top-right, units across
the middle, System title + rooms/ticks centred, then the premium-materials row.
Material tiles are the user's own tiles lifted from their System 4 design so the row matches.
"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image

PX = 9525
S = 1080
TEAL = RGBColor(0x5B, 0xAD, 0x92)
INK = RGBColor(0x11, 0x11, 0x11)
GREY = RGBColor(0x44, 0x44, 0x44)
FONT = "Arial"

SYSTEMS = {
    1: dict(out_h=300, in_w=296, sub="Single split · 1 room · 5 Ticks",
            indoor=[(8, 392)]),
    2: dict(out_h=320, in_w=282, sub="Up to 2 rooms · 5 Ticks",
            indoor=[(8, 338), (790, 338)]),
    3: dict(out_h=350, in_w=330, sub="Up to 3 rooms · 5 Ticks",
            indoor=[(8, 268), (8, 452), (742, 360)]),
}
BAND_TOP, BAND_H = 250, 400


def wh(path, w=None, h=None):
    iw, ih = Image.open(path).size
    return (w, round(w * ih / iw)) if w else (round(h * iw / ih), h)


def text(slide, x, y, w, h, s, size_px, *, bold=False, color=INK, align=PP_ALIGN.LEFT,
         line=None, name=None):
    box = slide.shapes.add_textbox(Emu(x * PX), Emu(y * PX), Emu(w * PX), Emu(h * PX))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    if line is not None:
        p.line_spacing = line
    run = p.add_run()
    run.text = s
    f = run.font
    f.name, f.size, f.bold, f.color.rgb = FONT, Pt(size_px * 0.75), bold, color
    if name:
        box.name = name
    return box


def picture(slide, path, x, y, w=None, h=None, name=None):
    ww, hh = wh(path, w=w, h=h)
    pic = slide.shapes.add_picture(path, Emu(round(x) * PX), Emu(round(y) * PX),
                                   Emu(ww * PX), Emu(hh * PX))
    if name:
        pic.name = name
    return pic


prs = Presentation()
prs.slide_width = Emu(S * PX)
prs.slide_height = Emu(S * PX)
blank = prs.slide_layouts[6]

for n, cfg in SYSTEMS.items():
    slide = prs.slides.add_slide(blank)

    picture(slide, "kool-logo.png", 18, 14, w=178, name="Kool logo")
    picture(slide, "starmex-logo.png", 222, 72, w=560, name="Mitsubishi Starmex logo")
    picture(slide, f"label-s{n}.png", 820, 74, h=136, name="NEA energy label")

    ow, oh = wh(f"unit-s{n}.png", h=cfg["out_h"])
    picture(slide, f"unit-s{n}.png", (S - ow) / 2, BAND_TOP + (BAND_H - oh) / 2, h=cfg["out_h"],
            name="Outdoor unit")
    for i, (x, y) in enumerate(cfg["indoor"], 1):
        picture(slide, "mit-indoor-fp.png" if n > 1 else "mit-indoor-gp.png", x, y, w=cfg["in_w"],
                name=f"Indoor unit {i}")

    text(slide, 0, 668, S, 110, f"System {n}", 86, bold=True, align=PP_ALIGN.CENTER,
         name="System title")
    text(slide, 0, 786, S, 40, cfg["sub"], 30, color=GREY, align=PP_ALIGN.CENTER, name="Subtitle")

    text(slide, 30, 888, 300, 140, "Free Upgrade To Premium Materials", 34, bold=True,
         color=TEAL, line=1.3, name="Free upgrade heading")
    for i, x in enumerate([349, 546, 753], 1):
        picture(slide, f"tile-{i}.png", x, 861, h=185, name=f"Material tile {i}")

prs.save("kool-shopee-covers-mitsubishi-s1-s3.pptx")
print("wrote kool-shopee-covers-mitsubishi-s1-s3.pptx")
