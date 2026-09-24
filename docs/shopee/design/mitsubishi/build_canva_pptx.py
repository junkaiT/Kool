"""Builds a 1024x1024 PowerPoint of the four Mitsubishi Starmex Shopee covers.

Canva imports .pptx and keeps text boxes and images as separate editable layers, so this
is the handover format for turning the cover into a Canva template.
Geometry mirrors the cover artboards in build_artboards.py (same 1024px square grid).
"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

PX = 9525  # EMU per CSS px at 96 dpi
S = 1024

TEAL = RGBColor(0x5B, 0xAD, 0x92)
INK = RGBColor(0x11, 0x11, 0x11)
GREY = RGBColor(0x55, 0x55, 0x55)
BORDER = RGBColor(0xE5, 0xE5, 0xE5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"  # safe everywhere; swap to Inter inside Canva to match the site

SYSTEMS = {
    1: dict(indoor="mit-indoor-gp.png", outdoor="mit-outdoor-s1.png", out_h=236, in_w=520, step=0, shift=0,
            sub="Single split · 1 room · 5 Ticks"),
    2: dict(indoor="mit-indoor-fp.png", outdoor="mit-outdoor-s2.png", out_h=244, in_w=460, step=118, shift=28,
            sub="Up to 2 rooms · 5 Ticks"),
    3: dict(indoor="mit-indoor-fp.png", outdoor="mit-outdoor-s3.png", out_h=276, in_w=476, step=104, shift=24,
            sub="Up to 3 rooms · 5 Ticks"),
    4: dict(indoor="mit-indoor-fp.png", outdoor="mit-outdoor-s4.png", out_h=290, in_w=400, step=84, shift=20,
            sub="Up to 4 rooms · 5 Ticks"),
}
TRUST = ["22g copper pipes", "K-Flex Titan", "Keystone branded cables"]


def wh(path, w=None, h=None):
    iw, ih = Image.open(path).size
    return (w, round(w * ih / iw)) if w else (round(h * iw / ih), h)


def text(slide, x, y, w, h, s, size_px, *, bold=False, color=INK, align=PP_ALIGN.LEFT,
         spacing=None, line=None, name=None):
    box = slide.shapes.add_textbox(Emu(x * PX), Emu(y * PX), Emu(w * PX), Emu(h * PX))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = s
    f = run.font
    f.name, f.size, f.bold, f.color.rgb = FONT, Pt(size_px * 0.75), bold, color
    if spacing is not None:
        rPr = run._r.get_or_add_rPr()
        rPr.set("spc", str(int(spacing * 0.75 * 100)))
    if line is not None:
        p.line_spacing = line
    if name:
        box.name = name
    return box


def picture(slide, path, x, y, w=None, h=None, name=None):
    ww, hh = wh(path, w=w, h=h)
    pic = slide.shapes.add_picture(path, Emu(x * PX), Emu(y * PX), Emu(ww * PX), Emu(hh * PX))
    if name:
        pic.name = name
    return pic


prs = Presentation()
prs.slide_width = Emu(S * PX)
prs.slide_height = Emu(S * PX)
blank = prs.slide_layouts[6]

for n, cfg in SYSTEMS.items():
    slide = prs.slides.add_slide(blank)

    # --- brand + title column (left) ---
    picture(slide, "starmex-logo.png", 52, 44, w=520, name="Mitsubishi Starmex logo")
    text(slide, 50, 206, 520, 150, f"System {n}", 120, bold=True, spacing=-3, line=1.0, name="System title")
    text(slide, 54, 348, 500, 40, cfg["sub"], 26, color=GREY, name="Subtitle")

    # --- materials card (right) ---
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(588 * PX), Emu(44 * PX),
                                  Emu(392 * PX), Emu(400 * PX))
    card.adjustments[0] = 0.04
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = BORDER
    card.line.width = Pt(1)
    card.shadow.inherit = False
    card.name = "Materials card"
    card.text_frame.text = ""
    picture(slide, "install-pipes-card.jpg", 608, 64, w=352, name="Materials photo")
    text(slide, 608, 230, 352, 70, "We only use high quality materials", 26, bold=True, line=1.15,
         name="Materials heading")
    for i, item in enumerate(TRUST):
        y = 312 + i * 34
        text(slide, 608, y, 26, 30, "✓", 24, bold=True, color=TEAL, name=f"Tick {i + 1}")
        text(slide, 638, y + 2, 320, 30, item, 21, bold=True, name=f"Material {i + 1}")

    # --- product photos ---
    for i in range(n):
        w = cfg["in_w"] * 1.155 if (n == 4 and i == 3) else cfg["in_w"]
        picture(slide, cfg["indoor"], 48 + i * cfg["shift"], 452 + i * cfg["step"], w=round(w),
                name=f"Indoor unit {i + 1}")
    ow, oh = wh(cfg["outdoor"], h=cfg["out_h"])
    picture(slide, cfg["outdoor"], S - 44 - ow, S - 24 - oh, w=ow, name="Outdoor unit")

    # --- brand sign-off ---
    picture(slide, "kool-logo.png", 52, 876, w=200, name="Kool logo")
    text(slide, 52, 972, 300, 30, "Kool & Kleen Aircon", 18, bold=True, color=GREY, name="Shop name")

prs.save("kool-shopee-covers-mitsubishi.pptx")
print("wrote kool-shopee-covers-mitsubishi.pptx —", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
